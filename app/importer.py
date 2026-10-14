from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import click
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select

from .domain import as_decimal
from .models import AreaSummary, CurvePeriod, Task

PHASE_KEYS = {
    "Curva S": "curva-s",
    "Curva S Preparação": "curva-s-preparacao",
    "Curva S Montagem": "curva-s-montagem",
    "Curva S Soldagem": "curva-s-soldagem",
    "Curva S Acabamento": "curva-s-acabamento",
    "Curva S Pintura": "curva-s-pintura",
}


def _date_value(value: Any, epoch: datetime) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    raise ValueError(f"Data inválida encontrada na planilha: {value!r}.")


def _datetime_value(value: Any, epoch: datetime) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted
        if isinstance(converted, date):
            return datetime.combine(converted, datetime.min.time())
    raise ValueError(f"Data e hora inválidas encontradas na planilha: {value!r}.")


def _number(value: Any, default: str = "0") -> Decimal:
    if value is None or value == "":
        value = default
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(
            "A planilha contém fórmula sem valor calculado em cache. "
            "Abra e salve o arquivo no Excel antes de importá-lo."
        )
    return as_decimal(value, "Peso na planilha")


def import_workbook(session, path: Path, replace: bool = False) -> dict[str, int]:
    if not path.is_file():
        raise FileNotFoundError(f"Arquivo de planilha não encontrado: {path}")

    occupied = any(
        session.scalar(select(model.id).limit(1)) is not None
        for model in (CurvePeriod, Task, AreaSummary)
    )
    if occupied and not replace:
        raise ValueError(
            "O banco já contém dados. Use --replace somente se deseja substituir "
            "as curvas, tarefas e resumos existentes."
        )

    workbook = load_workbook(path, data_only=True, read_only=True)
    missing = set(PHASE_KEYS) - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"Abas obrigatórias ausentes: {', '.join(sorted(missing))}.")

    counts = {"curves": 0, "tasks": 0, "areas": 0}
    try:
        if replace:
            session.query(CurvePeriod).delete(synchronize_session=False)
            session.query(Task).delete(synchronize_session=False)
            session.query(AreaSummary).delete(synchronize_session=False)

        for sheet_name, sheet_key in PHASE_KEYS.items():
            worksheet = workbook[sheet_name]
            start_row = 36
            for row_number in range(start_row, worksheet.max_row + 1):
                sequence, period_date = worksheet.cell(row_number, 1).value, worksheet.cell(
                    row_number, 2
                ).value
                planned = worksheet.cell(row_number, 3).value
                if sequence is None or period_date is None or planned is None:
                    continue
                session.add(
                    CurvePeriod(
                        sheet_key=sheet_key,
                        sequence=int(sequence),
                        period_date=_date_value(period_date, workbook.epoch),
                        planned_kg=_number(planned),
                        actual_kg=_number(worksheet.cell(row_number, 6).value),
                    )
                )
                counts["curves"] += 1

        tasks_sheet = workbook["Base Tarefas"]
        headers = {
            str(tasks_sheet.cell(1, column).value): column
            for column in range(1, tasks_sheet.max_column + 1)
            if tasks_sheet.cell(1, column).value is not None
        }
        required = {
            "ID MS Project",
            "Projeto/Área",
            "Desenho",
            "Marca/Item",
            "Nome da Tarefa",
            "Início",
            "Término",
            "Peso planejado (kg)",
            "Peso fabricado (kg)",
            "Peso expedido (kg)",
        }
        if not required.issubset(headers):
            raise ValueError("A aba Base Tarefas não contém as colunas esperadas.")
        for row_number in range(2, tasks_sheet.max_row + 1):
            source_id = tasks_sheet.cell(row_number, headers["ID MS Project"]).value
            name = tasks_sheet.cell(row_number, headers["Nome da Tarefa"]).value
            if source_id is None or not name:
                continue
            start_value = tasks_sheet.cell(row_number, headers["Início"]).value
            end_value = tasks_sheet.cell(row_number, headers["Término"]).value
            session.add(
                Task(
                    source_id=int(source_id),
                    project_area=str(
                        tasks_sheet.cell(row_number, headers["Projeto/Área"]).value or ""
                    ),
                    drawing=str(tasks_sheet.cell(row_number, headers["Desenho"]).value or ""),
                    item_code=str(tasks_sheet.cell(row_number, headers["Marca/Item"]).value or ""),
                    name=str(name),
                    starts_on=_datetime_value(start_value, workbook.epoch) if start_value else None,
                    ends_on=_datetime_value(end_value, workbook.epoch) if end_value else None,
                    planned_kg=_number(
                        tasks_sheet.cell(row_number, headers["Peso planejado (kg)"]).value
                    ),
                    manufactured_kg=_number(
                        tasks_sheet.cell(row_number, headers["Peso fabricado (kg)"]).value
                    ),
                    shipped_kg=_number(
                        tasks_sheet.cell(row_number, headers["Peso expedido (kg)"]).value
                    ),
                    source_row=row_number,
                    database_row=(
                        int(tasks_sheet.cell(row_number, headers["Linha Banco de Dados"]).value)
                        if "Linha Banco de Dados" in headers
                        and tasks_sheet.cell(row_number, headers["Linha Banco de Dados"]).value
                        is not None
                        else None
                    ),
                )
            )
            counts["tasks"] += 1

        summary_sheet = workbook["Resumo por Área"]
        for row_number in range(4, summary_sheet.max_row + 1):
            area = summary_sheet.cell(row_number, 1).value
            if not area:
                continue
            session.add(
                AreaSummary(
                    project_area=str(area),
                    task_count=int(summary_sheet.cell(row_number, 2).value or 0),
                    planned_kg=_number(summary_sheet.cell(row_number, 3).value),
                    manufactured_kg=_number(summary_sheet.cell(row_number, 4).value),
                )
            )
            counts["areas"] += 1

        session.flush()
        session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        workbook.close()
    return counts


def register_import_command(app) -> None:
    @app.cli.command("import-workbook")
    @click.argument("path", type=click.Path(exists=True, path_type=Path))
    @click.option("--replace", is_flag=True, help="Substitui os dados já importados.")
    def import_workbook_command(path: Path, replace: bool) -> None:
        factory = app.extensions["db_session_factory"]
        with factory() as session:
            counts = import_workbook(session, path, replace=replace)
        click.echo(
            f"Importação concluída: {counts['curves']} períodos de curva, "
            f"{counts['tasks']} tarefas e {counts['areas']} áreas."
        )
